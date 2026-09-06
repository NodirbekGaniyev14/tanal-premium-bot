"""Excel importi — Excel haqiqat manbai, baza qo'lda tahrirlanmaydi.

Har import avval to'liq tekshiriladi, keyin yoziladi: bitta xato qator bo'lsa
hech nima o'zgarmaydi va xabar aniq qator raqami bilan qaytadi.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import AiTask, AnswerKey, Level, SetModel, Student
from app.db.repo import get_or_create_cohort
from app.enums import AiKind, Section, StudentStatus
from app.services.answer_engine.levels import LevelBand, validate_bands
from app.utils.phone import normalize_phone

KINDS = ("sets", "answer_keys", "levels", "students", "ai_tasks")

_TRUTHY = {"yes", "ha", "true", "1", "y", "+"}
_FALSY = {"no", "yo'q", "yoq", "false", "0", "n", "-", ""}


@dataclass(slots=True)
class ImportReport:
    kind: str
    created: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def as_text(self) -> str:
        head = f"<b>{self.kind}</b> importi"
        if not self.ok:
            body = "\n".join(f"• {e}" for e in self.errors[:25])
            more = (
                f"\n… va yana {len(self.errors) - 25} ta xato"
                if len(self.errors) > 25
                else ""
            )
            return f"{head} — ❌ bekor qilindi.\n\n{body}{more}"
        lines = [f"{head} — ✅ tayyor.", f"Yangi: {self.created} · Yangilandi: {self.updated}"]
        if self.warnings:
            lines.append("")
            lines.extend(f"⚠️ {w}" for w in self.warnings[:15])
        return "\n".join(lines)


def detect_kind(file_name: str) -> str | None:
    name = (file_name or "").lower()
    for kind in ("ai_tasks", "answer_keys", "sets", "levels", "students"):
        if kind in name:
            return kind
    return None


def _read_rows(data: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return []
    columns = [str(c).strip().lower() if c is not None else "" for c in header]
    out: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        record = {columns[i]: row[i] for i in range(min(len(columns), len(row)))}
        record["_row"] = index
        out.append(record)
    workbook.close()
    return out


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(float(str(value).strip()))
    except ValueError:
        return None


def _bool(value: Any) -> bool | None:
    text = (str(value).strip().lower() if value is not None else "")
    if text in _TRUTHY:
        return True
    if text in _FALSY:
        return False
    return None


def _date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


# --- sets ------------------------------------------------------------------


async def import_sets(session: AsyncSession, data: bytes) -> ImportReport:
    report = ImportReport(kind="sets")
    rows = _read_rows(data)
    if not rows:
        report.errors.append("Fayl bo'sh yoki sarlavha qatori topilmadi.")
        return report

    valid_sections = {s.value for s in Section}
    seen: set[str] = set()
    prepared: list[dict[str, Any]] = []

    for row in rows:
        line = row["_row"]
        set_code = _text(row.get("set_code"))
        if not set_code:
            report.errors.append(f"sets, {line}-qator: set_code bo'sh")
            continue
        if set_code in seen:
            report.errors.append(f"sets, {line}-qator: {set_code} takrorlangan")
            continue
        seen.add(set_code)

        section = _text(row.get("section"))
        if section not in valid_sections:
            report.errors.append(
                f"sets, {line}-qator: section «{section}» noto'g'ri "
                f"(ruxsat: {', '.join(sorted(valid_sections))})"
            )
            continue

        total_q = _int(row.get("total_q"))
        if not total_q or total_q < 1:
            report.errors.append(f"sets, {line}-qator: total_q butun musbat son bo'lsin")
            continue

        attempts = _int(row.get("attempts")) or 1
        if attempts < 1:
            report.errors.append(f"sets, {line}-qator: attempts 1 dan kichik bo'lmasin")
            continue

        reveal = _bool(row.get("reveal"))
        if reveal is None:
            report.errors.append(f"sets, {line}-qator: reveal «yes» yoki «no» bo'lsin")
            continue

        title = _text(row.get("title"))
        if not title:
            report.errors.append(f"sets, {line}-qator: title bo'sh")
            continue

        prepared.append(
            {
                "set_code": set_code,
                "section": section,
                "title": title,
                "topic_no": _int(row.get("topic_no")),
                "total_q": total_q,
                "attempts_allowed": attempts,
                "reveal": reveal,
                "requires": _text(row.get("requires")),
                "sort_order": _int(row.get("sort_order")) or _int(row.get("topic_no")) or 0,
                "pdf_file": _text(row.get("pdf_file")),
                "audio_file": _text(row.get("audio_file")),
                "theory_file": _text(row.get("theory_file")),
            }
        )

    for item in prepared:
        if item["requires"] and item["requires"] not in seen:
            report.warnings.append(
                f"{item['set_code']}: requires «{item['requires']}» shu faylda yo'q"
            )

    if not report.ok:
        return report

    existing = {
        row.set_code: row
        for row in (await session.execute(select(SetModel))).scalars()
    }
    for item in prepared:
        current = existing.get(item["set_code"])
        if current is None:
            session.add(SetModel(**item))
            report.created += 1
            continue
        # Fayl nomi o'zgargan bo'lsa, eski `file_id` kuchini yo'qotadi.
        for name_column, id_column in (
            ("pdf_file", "pdf_file_id"),
            ("audio_file", "audio_file_id"),
            ("theory_file", "theory_file_id"),
        ):
            if getattr(current, name_column) != item[name_column]:
                setattr(current, id_column, None)
        for key, value in item.items():
            setattr(current, key, value)
        report.updated += 1

    await session.flush()
    return report


# --- answer_keys -----------------------------------------------------------


async def import_answer_keys(session: AsyncSession, data: bytes) -> ImportReport:
    report = ImportReport(kind="answer_keys")
    rows = _read_rows(data)
    if not rows:
        report.errors.append("Fayl bo'sh yoki sarlavha qatori topilmadi.")
        return report

    sets = {
        row.set_code: row for row in (await session.execute(select(SetModel))).scalars()
    }
    seen: dict[tuple[str, int], int] = {}
    prepared: list[dict[str, Any]] = []

    for row in rows:
        line = row["_row"]
        set_code = _text(row.get("set_code"))
        q_no = _int(row.get("q_no"))
        answer = _text(row.get("answer"))

        if not set_code:
            report.errors.append(f"answer_keys, {line}-qator: set_code bo'sh")
            continue
        if set_code not in sets:
            report.errors.append(
                f"answer_keys, {line}-qator: «{set_code}» sets ichida yo'q — "
                "avval sets.xlsx ni import qiling"
            )
            continue
        if not q_no or q_no < 1:
            report.errors.append(f"answer_keys, {line}-qator: q_no noto'g'ri")
            continue
        if q_no > sets[set_code].total_q:
            report.errors.append(
                f"answer_keys, {line}-qator: {set_code} da {q_no}-savol yo'q "
                f"(jami {sets[set_code].total_q})"
            )
            continue
        if not answer:
            report.errors.append(f"answer_keys, {line}-qator: answer bo'sh")
            continue
        key = (set_code, q_no)
        if key in seen:
            report.errors.append(
                f"answer_keys, {line}-qator: {set_code} uchun {q_no}-savol ikki marta "
                f"(birinchisi {seen[key]}-qatorda)"
            )
            continue
        seen[key] = line

        prepared.append(
            {
                "set_code": set_code,
                "q_no": q_no,
                "answer": answer,
                "accept_also": _text(row.get("accept_also")),
                "tag": _text(row.get("tag")),
            }
        )

    # uzilishlarni ogohlantirish sifatida ko'rsatamiz
    by_set: dict[str, set[int]] = {}
    for item in prepared:
        by_set.setdefault(item["set_code"], set()).add(item["q_no"])
    for set_code, numbers in by_set.items():
        expected = set(range(1, sets[set_code].total_q + 1))
        missing = sorted(expected - numbers)
        if missing:
            shown = ", ".join(str(n) for n in missing[:12])
            more = " …" if len(missing) > 12 else ""
            report.warnings.append(f"{set_code}: kalit yo'q savollar — {shown}{more}")
        if not any(item.get("tag") for item in prepared if item["set_code"] == set_code):
            report.warnings.append(
                f"{set_code}: tag ustuni bo'sh — tavsiyalar umumiy bo'ladi"
            )

    if not report.ok:
        return report

    existing = {
        (row.set_code, row.q_no): row
        for row in (
            await session.execute(
                select(AnswerKey).where(AnswerKey.set_code.in_(by_set.keys()))
            )
        ).scalars()
    }
    for item in prepared:
        current = existing.get((item["set_code"], item["q_no"]))
        if current is None:
            session.add(AnswerKey(**item))
            report.created += 1
        else:
            current.answer = item["answer"]
            current.accept_also = item["accept_also"]
            current.tag = item["tag"]
            report.updated += 1

    await session.flush()
    return report


# --- levels ----------------------------------------------------------------


async def import_levels(session: AsyncSession, data: bytes) -> ImportReport:
    report = ImportReport(kind="levels")
    rows = _read_rows(data)
    if not rows:
        report.errors.append("Fayl bo'sh yoki sarlavha qatori topilmadi.")
        return report

    prepared: list[dict[str, Any]] = []
    for row in rows:
        line = row["_row"]
        scope = _text(row.get("scope"))
        raw_min = _int(row.get("raw_min"))
        raw_max = _int(row.get("raw_max"))
        level = _text(row.get("level"))
        if not scope or raw_min is None or raw_max is None or not level:
            report.errors.append(
                f"levels, {line}-qator: scope, raw_min, raw_max, level to'ldirilsin"
            )
            continue
        prepared.append(
            {
                "scope": scope,
                "raw_min": raw_min,
                "raw_max": raw_max,
                "level": level,
                "label": _text(row.get("label")),
            }
        )

    by_scope: dict[str, list[LevelBand]] = {}
    for item in prepared:
        by_scope.setdefault(item["scope"], []).append(
            LevelBand(item["scope"], item["raw_min"], item["raw_max"], item["level"])
        )
    for scope, bands in by_scope.items():
        for problem in validate_bands(bands):
            report.errors.append(f"levels ({scope}): {problem}")

    if not report.ok:
        return report

    for scope in by_scope:
        for row in (
            await session.execute(select(Level).where(Level.scope == scope))
        ).scalars():
            await session.delete(row)
    await session.flush()

    for item in prepared:
        session.add(Level(**item))
        report.created += 1

    await session.flush()
    return report


# --- students --------------------------------------------------------------


async def import_students(session: AsyncSession, data: bytes) -> ImportReport:
    report = ImportReport(kind="students")
    rows = _read_rows(data)
    if not rows:
        report.errors.append("Fayl bo'sh yoki sarlavha qatori topilmadi.")
        return report

    seen: dict[str, int] = {}
    prepared: list[dict[str, Any]] = []

    for row in rows:
        line = row["_row"]
        full_name = _text(row.get("full_name"))
        phone = normalize_phone(_text(row.get("phone")))
        if not full_name:
            report.errors.append(f"students, {line}-qator: full_name bo'sh")
            continue
        if not phone:
            report.errors.append(
                f"students, {line}-qator: telefon raqamini o'qib bo'lmadi "
                f"«{row.get('phone')}»"
            )
            continue
        if phone in seen:
            report.errors.append(
                f"students, {line}-qator: {phone} takrorlangan "
                f"(birinchisi {seen[phone]}-qatorda)"
            )
            continue
        seen[phone] = line
        prepared.append(
            {
                "full_name": full_name,
                "phone": phone,
                "cohort": _text(row.get("cohort")),
                "note": _text(row.get("note")),
                "access_until": _date(row.get("access_until")),
            }
        )

    if not report.ok:
        return report

    cohorts: dict[str, int] = {}
    for item in prepared:
        name = item["cohort"]
        if name and name not in cohorts:
            cohorts[name] = (await get_or_create_cohort(session, name)).id

    existing = {
        row.phone: row
        for row in (
            await session.execute(select(Student).where(Student.phone.in_(seen.keys())))
        ).scalars()
    }

    for item in prepared:
        cohort_id = cohorts.get(item["cohort"]) if item["cohort"] else None
        current = existing.get(item["phone"])
        if current is None:
            session.add(
                Student(
                    full_name=item["full_name"],
                    phone=item["phone"],
                    cohort_id=cohort_id,
                    note=item["note"],
                    access_until=item["access_until"],
                    status=StudentStatus.INVITED,
                )
            )
            report.created += 1
            continue
        # tg_id va status tegilmaydi — o'quvchi allaqachon kirgan bo'lishi mumkin
        current.full_name = item["full_name"]
        if cohort_id:
            current.cohort_id = cohort_id
        if item["note"]:
            current.note = item["note"]
        if item["access_until"]:
            current.access_until = item["access_until"]
        report.updated += 1

    await session.flush()
    return report


# --- ai_tasks --------------------------------------------------------------

#: Har bo'lim uchun ruxsat etilgan `part` qiymatlari.
_AI_PARTS = {
    AiKind.WRITING.value: {"task1", "task2"},
    AiKind.SPEAKING.value: {"1", "2", "3"},
}


async def import_ai_tasks(session: AsyncSession, data: bytes) -> ImportReport:
    report = ImportReport(kind="ai_tasks")
    rows = _read_rows(data)
    if not rows:
        report.errors.append("Fayl bo'sh yoki sarlavha qatori topilmadi.")
        return report

    seen: set[str] = set()
    prepared: list[dict[str, Any]] = []

    for row in rows:
        line = row["_row"]
        code = _text(row.get("code"))
        kind = (_text(row.get("kind")) or "").lower()
        part = _text(row.get("part"))
        title = _text(row.get("title"))
        prompt_text = _text(row.get("prompt_text"))

        if not code:
            report.errors.append(f"ai_tasks, {line}-qator: code bo'sh")
            continue
        if code in seen:
            report.errors.append(f"ai_tasks, {line}-qator: {code} takrorlangan")
            continue
        seen.add(code)

        if kind not in _AI_PARTS:
            report.errors.append(
                f"ai_tasks, {line}-qator: kind «{kind}» noto'g'ri "
                "(writing yoki speaking)"
            )
            continue
        if part is not None:
            part = str(part).strip().lower().replace(" ", "")
        if part not in _AI_PARTS[kind]:
            report.errors.append(
                f"ai_tasks, {line}-qator: {kind} uchun part «{part}» noto'g'ri "
                f"(ruxsat: {', '.join(sorted(_AI_PARTS[kind]))})"
            )
            continue
        if not title:
            report.errors.append(f"ai_tasks, {line}-qator: title bo'sh")
            continue
        if not prompt_text:
            report.errors.append(f"ai_tasks, {line}-qator: prompt_text bo'sh")
            continue

        prepared.append(
            {
                "code": code,
                "kind": kind,
                "part": part,
                "title": title,
                "prompt_text": prompt_text,
                "cue_points": _text(row.get("cue_points")),
                "image_file": _text(row.get("image_file")),
            }
        )

    for item in prepared:
        speaking_part_2 = (
            item["kind"] == AiKind.SPEAKING.value and item["part"] == "2"
        )
        if speaking_part_2 and not item["cue_points"]:
            report.warnings.append(f"{item['code']}: Part 2 uchun cue_points bo'sh")

    if not report.ok:
        return report

    existing = {
        row.code: row
        for row in (
            await session.execute(select(AiTask).where(AiTask.code.in_(seen)))
        ).scalars()
    }
    for item in prepared:
        current = existing.get(item["code"])
        if current is None:
            session.add(AiTask(**item))
            report.created += 1
            continue
        if current.image_file != item["image_file"]:
            current.image_file_id = None
        for key, value in item.items():
            setattr(current, key, value)
        report.updated += 1

    await session.flush()
    return report


IMPORTERS = {
    "sets": import_sets,
    "answer_keys": import_answer_keys,
    "levels": import_levels,
    "students": import_students,
    "ai_tasks": import_ai_tasks,
}


async def run_import(session: AsyncSession, kind: str, data: bytes) -> ImportReport:
    handler = IMPORTERS.get(kind)
    if handler is None:
        return ImportReport(kind=kind, errors=[f"Noma'lum import turi: {kind}"])
    return await handler(session, data)
